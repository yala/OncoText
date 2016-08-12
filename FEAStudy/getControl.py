import pickle as p
import sys
sys.path.append("..")
from constants import *
from helpers import *
from csv import DictWriter
import xmltodict
import openpyxl as xl



diagnoses = [ u'DCIS', u'ILC', u'Phylodes', u'CancerSuspicious', \
 u'CancerNOS', u'TubularCancer', u'IDC', u'LCIS', 'InsituNOS', \
 u'ADH', u'CancerInvInsituNOS',\
 u'ClingingCarcinoma', u'InvasiveSuspicious', u'LobularNeoplasia', 
 u'Medullary', u'ALH', u'ADH_DCIS', 'Borderline', 'FlatEpithelial']


#Return true if the report is benign
def filt(report):
    if report['DOBPlusX'] == 0:
        return False
    for diag in diagnoses:
            assert report[diag] == 1 or report[diag] == 0
            if report[diag] == 1:
                return False
    return True

def makeAlreadyExportedDict(feas, possControls):
    visited = set()
    controlled = set()
    feaMRNs = [f['MRNPlusX'] for f in feas]
    wb = xl.load_workbook('FEAControlGroupFromAdam_8.2.2016.xlsx')
    sheet = wb.get_active_sheet()
    legend = []
    legendBool = True
    for row in sheet.iter_rows():
        if legendBool:
            for indx, col in enumerate(row):
                legend.append(col.value)
            legendBool = False
            print "Legend is", legend
        else:
            rowDict = {}
            for indx, col in enumerate(row):
                ## Use makes mrns into strings and leave dates alone
                rowDict[legend[indx]] = str(col.value) if indx == 0 else col.value

            visited.add(str(rowDict['MRNPlusX']))
            controlled.add(str(rowDict['ControlForMRNActualPatientInStudyFEA+x']))
    assert '2069287' in controlled
    assert '20913633' in controlled

    return visited, controlled

def getExtraControls():
    mrnToExtraDict= {}
    wb = xl.load_workbook('NeedMoreControlsForTheseMRNsAdam.xlsx')
    sheet = wb.get_active_sheet()
    legend = []
    legendBool = True
    for row in sheet.iter_rows():
        if legendBool:
            for indx, col in enumerate(row):
                legend.append(col.value)
            legendBool = False
            print "Legend is", legend
        else:
            rowDict = {}
            for indx, col in enumerate(row):
                ## Use makes mrns into strings and leave dates alone
                rowDict[legend[indx]] = str(col.value) if indx == 0 else col.value
            mrn = str(rowDict['MRNCaseFEAStudyPlusX'])
            num = rowDict['NumberNewControlsNeeded']

            mrnToExtraDict[mrn] = num
    return mrnToExtraDict

### Main Script
### This script opens up first records from the FEA patients, and 
### tries to find maxControl number patients to match each one
### Control patient has to have benign episode, be within an age Tol of 
### age, and within a monthTol months of report date

feas = p.load(open('FirstRecordsForFEAs.p', 'rb'))
mrnDict = getMRNDB(p.load(open('../'+outPath, 'rb')))
oldDobs1 = xmltodict.parse(open('DOBmrnplusx.xml', 'rb'))['dataroot']['DOBmrnplusx']
oldDobs2 = xmltodict.parse( open('DataForAdamOn2010_2016WithDOB.xml', 'rb') )['dataroot']['DataForAdamOn2010_2016']
newDobs = p.load(open('graceLatestDOBs.p', 'rb'))   

print "Len FEAS", len(feas)
feaDOBs = [f for f in feas if not f['DOBPlusX'] == 0]
feaNoDOBs = [f['MRNPlusX'] for f in feas if  f['DOBPlusX'] == 0]
feaMRNs = [f['MRNPlusX'] for f in feaDOBs]
print 'FEA NO DOBS', feaNoDOBs
print "Len FEAS DOBs", len(feaDOBs)

possControls = []
for mrn in mrnDict:
    reports = mrnDict[mrn]
    epId = reports[0]['EpisodeID']
    ## Comperehension asserts that all reports in episode of first report is benign
    if not False in [filt(r) for r in reports if r['EpisodeID'] == epId]:
        possControls.append(mrn)

print "Len Poss Controls", len(possControls)

ageTol = 2
maxControl = 4
monthTol = 3

numPosExamples = 0
i = 0

numNegExamples = 0
j= 0

getNewestBool = True
visited = set()
controlled = set()

if getNewestBool:
    visited, controlled = makeAlreadyExportedDict(feaDOBs, possControls)
    extraExtracts = getExtraControls()

    print "Len visited in prev set", len(visited)
    for extra in extraExtracts:
        assert extra in controlled
        assert extra in feaMRNs

## This has number of controls we need per MRN, lets us get extra controls selectively
## For things that are controlled, that sould be 0
maxControlDict = {}


for mrn in feaMRNs:
    if getNewestBool and mrn in controlled:
        if mrn in extraExtracts:
            maxControlDict[mrn] = extraExtracts[mrn]
        else:
            maxControlDict[mrn] = 0
    else:
        maxControlDict[mrn] = maxControl

countNoControl = 0
controlListMRNs = []
controlListRecords = []
notEnoughMRNs = []
## For each fea, find certain number of controls
for fea in feaDOBs:
    dob = fea['DOBPlusX']
    diagDate = fea['Report_Date_TimePlusX']
    feaMRN =  fea['MRNPlusX']
    controls = 0
    #Find controls for each one
    for mrn in possControls:
        if mrn in visited:# or not 'breast' in mrnDict[mrn][0]['Report_Text'].lower(): #Make sure we haven't used this control before
            continue
        possDob = mrnDict[mrn][0]['DOBPlusX']
        possDiagDate = mrnDict[mrn][0]['Report_Date_TimePlusX']

        ageDif = abs((possDob - dob).days)/365.
        diagDif = abs((possDiagDate - diagDate).days)/30.

        #Control checks
        if ageDif <= ageTol and controls < maxControlDict[feaMRN] and \
        diagDif <= monthTol:
            controlListMRNs.append( {'MRNPlusX': mrn , \
            'Institution': mrnDict[mrn][0]['Institution'],\
            'ControlForMRN': feaMRN, 'ControlForInst': fea['Institution']   })
            
            for rep in mrnDict[mrn]:
                rep['ControlForMRN'] = feaMRN
                rep['ControlForInst'] = fea['Institution']

            controlListRecords.extend(mrnDict[mrn])
            controls += 1
            visited.add(mrn)

            if i < numPosExamples:
                print "POS"
                print "DOBs", dob,  possDob, ageDif
                print "diag", diagDate, possDiagDate, diagDif
                i += 1
        else:
            if j < numNegExamples:
                print "NEG"
                print "DOBs", dob,  possDob, ageDif
                print "diag", diagDate, possDiagDate, diagDif
                j += 1

    print dob.year, "Controls: ", controls

    if controls < maxControlDict[feaMRN]:
        countNoControl += 1
        notEnoughMRNs.append(feaMRN)



print "Len FEAS DOBs", len(feaDOBs)
print "Count FEA mrns with not enough control", countNoControl
print "Those MRNs are", notEnoughMRNs
print "Num mrns in control set", len(controlListMRNs)
print "Num elts in control set", len(controlListRecords)

print "Controls cover", len(set([(r['ControlForMRN'], r['ControlForInst']) for r in controlListRecords]))

writeSimple = False
if writeSimple:
    mrnKeys = ['MRNPlusX', 'Institution', 'ControlForMRN', 'ControlForInst']
    with open("controlGroupMRN_Inst_ControlFor.csv", "wb") as outCSVMrn:
        writer = DictWriter(outCSVMrn, mrnKeys)
        writer.writeheader()
        writer.writerows(controlListMRNs)

writeAll = True
writeAllPath = "controlGroupAllData8_11.csv" if getNewestBool else "controlGroupAllData.csv"  
if writeAll:
    keys = controlListRecords[0].keys()
    with open(writeAllPath, "wb") as outCSVRecords:
        writer = DictWriter(outCSVRecords, keys)
        writer.writeheader()
        writer.writerows(controlListRecords)

missingDOBFeas = [f['MRNPlusX'] for f in feas if f['DOBPlusX'] == 0]
print "Number of missing DOBs", len(missingDOBFeas)
with open('missingFEAMRNs.csv', 'wb') as outMissing:
    outMissing.write(','.join(missingDOBFeas))



            






